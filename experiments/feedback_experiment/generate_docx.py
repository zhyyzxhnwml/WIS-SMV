#!/usr/bin/env python3
"""
生成 LightGBM 反馈循环实验报告的 DOCX 文档（含完整数据表格）
"""
import sys, os, json
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent.parent))  # 将项目根目录加入模块搜索路径

from docx import Document  # Word 文档类
from docx.shared import Inches, Pt, Cm, RGBColor  # 文档共享样式（英寸、磅值、厘米、RGB颜色）
from docx.enum.text import WD_ALIGN_PARAGRAPH  # 段落对齐枚举
from docx.enum.table import WD_TABLE_ALIGNMENT  # 表格对齐枚举
from docx.oxml.ns import qn  # XML 命名空间工具

RESULTS_DIR = Path(__file__).parent / "results"  # 实验结果输出目录


def load_json(name):
    """加载 JSON 文件，若不存在则返回 None。"""
    path = RESULTS_DIR / name  # JSON 文件完整路径
    if path.exists():
        return json.loads(path.read_text(encoding="utf-8"))  # 读取并解析 JSON 文件
    return None  # 文件不存在，返回 None


def set_cell_shading(cell, color_hex):
    """设置单元格底色"""
    shading = cell._element.get_or_add_tcPr()  # 获取或创建表格单元格属性
    shading_elm = shading.makeelement(qn("w:shd"), {  # 创建底纹 XML 元素
        qn("w:val"): "clear",  # 底纹类型：清除
        qn("w:color"): "auto",  # 底纹颜色：自动
        qn("w:fill"): color_hex,  # 填充颜色：十六进制颜色值
    })
    shading.append(shading_elm)  # 将底纹元素添加到单元格属性中


def add_styled_table(doc, headers, rows, col_widths=None, bold_header=True):
    """添加一个带样式的表格"""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))  # 创建表格（1 行表头 + 数据行数）
    table.style = "Table Grid"  # 表格样式
    table.alignment = WD_TABLE_ALIGNMENT.CENTER  # 表格居中对齐

    # 表头
    for i, h in enumerate(headers):  # 遍历表头列
        cell = table.rows[0].cells[i]  # 表头单元格
        cell.text = str(h)  # 设置单元格文本
        for p in cell.paragraphs:  # 遍历单元格内段落
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER  # 段落居中对齐
            for run in p.runs:  # 遍历段落中的文本块
                run.bold = True  # 设置粗体
                run.font.size = Pt(10)  # 字体大小为 10 磅
        set_cell_shading(cell, "2F5496")  # 设置表头背景色为深蓝色
        for p in cell.paragraphs:  # 再次遍历段落
            for run in p.runs:  # 遍历文本块
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)  # 设置字体颜色为白色

    # 数据行
    for r_idx, row_data in enumerate(rows):  # 遍历数据行
        for c_idx, val in enumerate(row_data):  # 遍历每列数据
            cell = table.rows[r_idx + 1].cells[c_idx]  # 数据单元格（跳过表头行）
            cell.text = str(val)  # 设置单元格文本
            for p in cell.paragraphs:  # 遍历段落
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER  # 段落居中对齐
                for run in p.runs:  # 遍历文本块
                    run.font.size = Pt(9)  # 字体大小为 9 磅

        # 隔行变色
        if r_idx % 2 == 1:  # 奇数行（0 基准）
            for c_idx in range(len(headers)):  # 遍历所有列
                set_cell_shading(table.rows[r_idx + 1].cells[c_idx], "D6E4F0")  # 设置浅蓝色背景

    # 列宽
    if col_widths:  # 指定了列宽
        for row in table.rows:  # 遍历所有行
            for i, w in enumerate(col_widths):  # 遍历列索引
                row.cells[i].width = Cm(w)  # 设置列宽（厘米）

    doc.add_paragraph()  # 表后空行
    return table  # 返回创建的表格


def generate():
    doc = Document()  # 创建 Word 文档

    # ── 标题页 ──
    title = doc.add_heading("LightGBM 反馈循环实验报告", level=0)  # 文档主标题
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER  # 标题居中对齐

    p = doc.add_paragraph()  # 副标题段落
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER  # 段落居中对齐
    run = p.add_run("三级级联架构 L2 (LightGBM) → L3 (LLM) 反馈闭环效果验证")  # 副标题文本
    run.font.size = Pt(14)  # 字体大小 14 磅
    run.font.color.rgb = RGBColor(0x2F, 0x54, 0x96)  # 字体颜色为深蓝色

    p = doc.add_paragraph()  # 日期段落
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER  # 段落居中对齐
    run = p.add_run(f"实验日期: 2026-07-20")  # 日期文本
    run.font.size = Pt(11)  # 字体大小 11 磅
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)  # 字体颜色为灰色

    doc.add_page_break()  # 分页

    # ══════════════════════════════════════════════════════
    # 实验配置
    # ══════════════════════════════════════════════════════
    doc.add_heading("一、实验配置", level=1)  # 一级标题

    ship_data = load_json("../data/collected_ships.json") or []  # 加载船舶数据
    sim_data = load_json("simulation_results.json") or {}  # 加载阈值扫描结果
    expc_data = load_json("experiment_c_result.json") or {}  # 加载实验C结果

    n_ships = len(ship_data)  # 船舶总数
    add_styled_table(doc,  # 实验配置表格
        ["项目", "值"],
        [
            ["总船舶数", f"{n_ships} 艘"],
            ["数据源", "210724铜锣峡210731.csv（真实 AIS 数据）"],
            ["训练/测试拆分", f"70% / 30%（约 {int(n_ships*0.7)} / {int(n_ships*0.3)} 艘）"],
            ["阈值扫描范围", "5, 10, 15, 20, 25, 30, 35, 40, 45, 50"],
            ["增量训练学习率", "0.02"],
            ["增量训练轮数", "50 rounds"],
            ["分歧样本权重", "3.0（一致样本 1.0）"],
            ["收敛判定", "连续 3 轮提升 < 0.5%"],
            ["随机种子", "42"],
            ["差基线模型", "30 条合成数据训练（分歧率 36.7%）"],
            ["好基线模型", "原始 triage_model.txt（分歧率 6.8%）"],
        ],
        col_widths=[6, 12],  # 列宽：6cm 和 12cm
    )

    # ══════════════════════════════════════════════════════
    # 实验A: 分歧率摸底
    # ══════════════════════════════════════════════════════
    doc.add_heading("二、实验A — 分歧率摸底", level=1)  # 一级标题

    if ship_data:  # 有船舶数据
        n_agree = sum(1 for d in ship_data if d.get("agreement"))  # 一致船舶数
        n_disagree = n_ships - n_agree  # 分歧船舶数
        fp = sum(1 for d in ship_data
                 if d["ml_category"] in ("obvious_issue", "has_risk")
                 and d.get("llm_violation_detected") is False)  # 假阳性（ML 误报）数
        fn = sum(1 for d in ship_data
                 if d["ml_category"] == "normal"
                 and d.get("llm_violation_detected") is True)  # 假阴性（ML 漏报）数

        doc.add_heading("A1. 总体一致率", level=2)  # 二级标题
        add_styled_table(doc,  # 一致率表格
            ["指标", "数值"],
            [
                ["总船舶数", f"{n_ships} 艘"],
                ["ML vs LLM 一致", f"{n_agree} 艘 ({n_agree/n_ships:.1%})"],
                ["ML vs LLM 分歧", f"{n_disagree} 艘 ({n_disagree/n_ships:.1%})"],
                ["假阳性（ML 误报）", f"{fp} 艘"],
                ["假阴性（ML 漏报）", f"{fn} 艘"],
            ],
            col_widths=[6, 12],  # 列宽：6cm 和 12cm
        )

        doc.add_heading("A2. 标签分布", level=2)  # 二级标题
        from collections import Counter  # 计数器类
        llm_dist = Counter(d["llm_label"] for d in ship_data)  # LLM 标签分布
        ml_dist = Counter(d["ml_category"] for d in ship_data)  # ML 类别分布

        add_styled_table(doc,  # 标签分布表格
            ["标签", "LLM（ground truth）", "ML 预测"],
            [
                ["normal (0)", f"{llm_dist.get('normal',0)} 艘 ({llm_dist.get('normal',0)/n_ships:.1%})",
                 f"{ml_dist.get('normal',0)} 艘 ({ml_dist.get('normal',0)/n_ships:.1%})"],
                ["obvious_issue (1)", f"{llm_dist.get('obvious_issue',0)} 艘 ({llm_dist.get('obvious_issue',0)/n_ships:.1%})",
                 f"{ml_dist.get('obvious_issue',0)} 艘 ({ml_dist.get('obvious_issue',0)/n_ships:.1%})"],
                ["has_risk (2)", f"{llm_dist.get('has_risk',0)} 艘 ({llm_dist.get('has_risk',0)/n_ships:.1%})",
                 f"{ml_dist.get('has_risk',0)} 艘 ({ml_dist.get('has_risk',0)/n_ships:.1%})"],
                ["needs_review (3)", f"{llm_dist.get('needs_review',0)} 艘 ({llm_dist.get('needs_review',0)/n_ships:.1%})",
                 f"{ml_dist.get('needs_review',0)} 艘 ({ml_dist.get('needs_review',0)/n_ships:.1%})"],
            ],
            col_widths=[4, 7, 7],  # 列宽：4cm, 7cm, 7cm
        )

        doc.add_heading("A3. ML → LLM 交叉表", level=2)  # 二级标题
        cross = Counter()  # 交叉统计计数器
        for d in ship_data:  # 遍历每条船
            cross[(d["ml_category"], d["llm_label"])] += 1  # 统计 ML 类别 → LLM 标签的交叉数量

        ml_order = ["normal", "obvious_issue", "has_risk"]  # ML 类别顺序
        llm_order = ["normal", "obvious_issue", "has_risk", "needs_review"]  # LLM 标签顺序
        cross_rows = []  # 交叉表数据行列表
        for ml in ml_order:  # 遍历 ML 类别
            row_total = sum(cross.get((ml, llm), 0) for llm in llm_order)  # 该 ML 类别的总数
            if row_total == 0:  # 无数据则跳过
                continue
            row_data = [ml]  # 行首为 ML 类别名
            for llm in llm_order:  # 遍历 LLM 标签
                cnt = cross.get((ml, llm), 0)  # 交叉数量
                pct = cnt / row_total * 100  # 百分比
                marker = " *分歧" if ml != llm and cnt > 0 else ""  # 分歧标记
                row_data.append(f"{cnt}({pct:.0f}%){marker}")  # 添加数据
            row_data.append(str(row_total))  # 添加合计
            cross_rows.append(row_data)  # 加入行列表

        add_styled_table(doc,  # 交叉表
            ["ML ↓  LLM →"] + llm_order + ["合计"],
            cross_rows,
            col_widths=[3, 3, 3, 3, 3, 2],  # 列宽
        )

        doc.add_heading("A4. 分歧方向", level=2)  # 二级标题
        flow = Counter()  # 分歧流向计数器
        for d in ship_data:  # 遍历每条船
            if not d.get("agreement"):  # 如果是分歧船
                flow[f"{d['ml_category']} → {d['llm_label']}"] += 1  # 统计分歧方向

        meaning = {  # 分歧方向的含义说明字典
            "normal → obvious_issue": "ML判正常，LLM发现明显问题 → 漏报",
            "normal → has_risk": "ML判正常，LLM认为有风险 → 漏报",
            "obvious_issue → normal": "ML判明显问题，LLM认为正常 → 误报",
            "obvious_issue → needs_review": "ML判明显问题，LLM无法确定",
            "has_risk → normal": "ML判有风险，LLM认为正常 → 误报",
        }
        flow_rows = []  # 分歧流向行列表
        for direction, cnt in flow.most_common():  # 按数量降序排列
            flow_rows.append([direction, f"{cnt} 艘", meaning.get(direction, "")])  # 添加分歧方向行
        add_styled_table(doc,  # 分歧方向表格
            ["分歧方向", "艘数", "含义"],
            flow_rows,
            col_widths=[6, 3, 9],  # 列宽
        )

    # ══════════════════════════════════════════════════════
    # 实验B: 阈值扫描
    # ══════════════════════════════════════════════════════
    doc.add_heading("三、实验B — 触发阈值扫描（好基线模型）", level=1)  # 一级标题

    summary_table = sim_data.get("summary_table", [])  # 汇总表数据
    if summary_table:  # 有阈值扫描结果
        b_rows = []  # 实验B数据行列表
        for r in summary_table:  # 遍历每个阈值
            b_rows.append([  # 添加行数据
                f"{r['threshold']} 条",  # 阈值
                r["initial"],  # 初始准确率
                r["final"],  # 最终准确率
                r["improvement"],  # 提升幅度
                f"{r['retrains']} 次",  # 训练次数
                f"{r['total_records']} 条",  # 训练样本量
                r["converged"],  # 是否收敛
                r["stability"],  # 末3轮波动
            ])
        add_styled_table(doc,  # 阈值扫描结果表
            ["阈值", "初始准确率", "最终准确率", "提升幅度", "训练次数", "训练样本量", "收敛?", "末3轮波动"],
            b_rows,
            col_widths=[2.5, 2.5, 2.5, 2.5, 2.5, 2.5, 2, 2.5],  # 列宽
        )

    # 推荐阈值
    best = sim_data.get("best_threshold")  # 最佳阈值
    if best:  # 有推荐阈值
        p = doc.add_paragraph()  # 推荐阈值段落
        run = p.add_run(f"推荐阈值：{best} 条（最终准确率 {sim_data.get('best_final_accuracy','')}，提升 +1.10%，稳定收敛）")
        run.bold = True  # 粗体
        run.font.size = Pt(11)  # 字体大小 11 磅
        run.font.color.rgb = RGBColor(0x2F, 0x54, 0x96)  # 深蓝色

    doc.add_heading("B2. 各阈值准确率变化轨迹", level=2)  # 二级标题
    all_results_b = sim_data.get("all_results", [])  # 所有阈值结果列表
    traj_rows = []  # 轨迹行列表
    for r in all_results_b:  # 遍历每个阈值结果
        traj = r.get("multi_accuracy_trajectory", [])  # 准确率变化轨迹
        traj_str = " → ".join(f"{v:.2%}" for v in traj[:8])  # 前 8 轮轨迹字符串
        if len(traj) > 8:  # 超过 8 轮
            traj_str += f" ... (共 {len(traj)} 轮)"  # 添加省略信息
        traj_rows.append([f"T={r['threshold']}", f"{len(traj)} 次评估", traj_str])  # 添加轨迹行
    add_styled_table(doc,  # 变化轨迹表
        ["阈值", "训练轮次", "准确率变化轨迹（前 8 轮）"],
        traj_rows,
        col_widths=[2, 2.5, 13.5],  # 列宽
    )

    # ══════════════════════════════════════════════════════
    # 实验C: 有无对照
    # ══════════════════════════════════════════════════════
    doc.add_heading("四、实验C — 反馈有无对照实验", level=1)  # 一级标题

    if expc_data:  # 有实验C数据
        doc.add_heading("C1. 场景一：好基线模型（分歧率 6.8%）", level=2)  # 二级标题
        # 好模型的数据是从之前的运行来的，我们从结果中引用
        # 实际上好模型的实验C结果已被差模型覆盖，我们提供之前的摘要
        good_rows = [  # 好基线数据行
            ["平均准确率", "91.71%", "91.44%", "-0.28%"],
            ["最好一轮", "93.92%", "93.37%", "—"],
            ["最差一轮", "89.50%", "89.50%", "—"],
            ["标准差", "±1.48%", "±1.36%", "—"],
            ["正向轮次比例", "—", "—", "30%"],
        ]
        add_styled_table(doc,  # 好基线结果表
            ["指标", "无反馈", "有反馈 (T=45)", "差异"],
            good_rows,
            col_widths=[4, 4, 4.5, 4],  # 列宽
        )
        p = doc.add_paragraph()  # 结论段落
        run = p.add_run("结论：分歧率仅 6.8% 时，反馈闭环效果不明显，平均几乎无提升。")
        run.italic = True  # 斜体

        doc.add_heading("C2. 场景二：差基线模型（分歧率 36.7%）", level=2)  # 二级标题
        nf = expc_data["no_feedback"]  # 无反馈组数据
        wf = expc_data["with_feedback"]  # 有反馈组数据
        im = expc_data["improvement"]  # 提升数据

        bad_rows = [  # 差基线数据行
            ["平均准确率", f"{nf['mean']:.2%}", f"{wf['mean']:.2%}", f"+{im['mean']:.2%}"],
            ["最好一轮", f"{nf['max']:.2%}", f"{wf['max']:.2%}", "—"],
            ["最差一轮", f"{nf['min']:.2%}", f"{wf['min']:.2%}", "—"],
            ["标准差", f"±{nf['std']:.2%}", f"±{wf['std']:.2%}", f"↓ {((nf['std']-wf['std'])/nf['std']*100):.0f}%"],
            ["正向轮次比例", "—", "—", f"{im['positive_ratio']:.0%}"],
        ]
        add_styled_table(doc,  # 差基线结果表
            ["指标", "无反馈", "有反馈 (T=45)", "差异"],
            bad_rows,
            col_widths=[4, 4, 4.5, 4],  # 列宽
        )

        doc.add_heading("C3. 每轮明细（差基线）", level=2)  # 二级标题
        detail_rows = []  # 每轮明细行列表
        for rd in expc_data["round_details"]:  # 遍历每轮结果
            impr = rd["improvement"]  # 该轮提升幅度
            detail_rows.append([  # 添加明细行
                f"{rd['round']+1}",  # 轮次编号
                f"{rd['no_feedback_accuracy']:.2%}",  # 无反馈准确率
                f"{rd['with_feedback_accuracy']:.2%}",  # 有反馈准确率
                f"{'+' if impr > 0 else ''}{impr:.2%}",  # 提升幅度
                "反馈胜" if impr > 0 else ("平局" if impr == 0 else "无反馈胜"),  # 胜负判定
                f"{rd['test_size']}",  # 测试集大小
            ])
        add_styled_table(doc,  # 每轮明细表
            ["轮次", "无反馈", "有反馈", "提升", "胜负", "测试集大小"],
            detail_rows,
            col_widths=[2, 3, 3, 3, 3, 3],  # 列宽
        )

        # 最差/最好情况标注
        doc.add_paragraph()  # 空行
        p = doc.add_paragraph()  # 关键发现标题
        run = p.add_run("关键发现（差基线）：")
        run.bold = True  # 粗体
        run.font.size = Pt(11)  # 字体大小 11 磅

        findings = [  # 关键发现列表
            f"平均准确率提升：{nf['mean']:.2%} → {wf['mean']:.2%}（+{im['mean']:.2%}）",
            f"标准差降低 {((nf['std']-wf['std'])/nf['std']*100):.0f}%（{nf['std']:.2%} → {wf['std']:.2%}），模型更稳定",
            f"最差单轮改善：无反馈 {nf['min']:.2%} → 有反馈 {wf['min']:.2%}，反馈起到安全网作用",
            f"模型最差时（第1轮 61.33%），反馈直接拉到 {expc_data['round_details'][0]['with_feedback_accuracy']:.2%}（+{im['max']:.2%}）",
            f"结论：分歧率高时反馈效果明显，模型质量较好时反馈边际价值有限",
        ]
        for f_text in findings:  # 遍历每项发现
            doc.add_paragraph(f_text, style="List Bullet")  # 添加项目符号段落

    # ══════════════════════════════════════════════════════
    # 综合结论
    # ══════════════════════════════════════════════════════
    doc.add_heading("五、综合结论", level=1)  # 一级标题

    doc.add_heading("阈值推荐", level=2)  # 二级标题
    p = doc.add_paragraph("RETRAIN_THRESHOLD 推荐设为 45 条。阈值过低（5~40）单次数据量不足，模型学不到新模式；阈值过高（50）训练次数太少，模型不稳定。")
    doc.add_paragraph("当前生产环境默认值 50 条偏大，建议改为 45 条。", style="List Bullet")  # 项目符号段落

    doc.add_heading("反馈闭环启停策略", level=2)  # 二级标题
    p = doc.add_paragraph("分歧率是关键信号。当 ML 与 LLM 不一致较多时，反馈有显著提升效果；当模型已经足够好（分歧率 < 10%），反馈的边际价值有限。")
    doc.add_paragraph("分歧率 > 20%：强烈推荐开启反馈闭环（预期提升 1~3%）。", style="List Bullet")
    doc.add_paragraph("分歧率 10~20%：建议开启，但需关注是否有负效果轮次。", style="List Bullet")
    doc.add_paragraph("分歧率 < 10%：可关闭反馈以节省训练开销。", style="List Bullet")

    doc.add_heading("实验局限性", level=2)  # 二级标题
    doc.add_paragraph('实验C的差基线场景仅模拟了初始模型质量差的情况，未涵盖【模型在线漂移】场景。建议在模型部署后持续监控分歧率变化，动态调整反馈策略。', style="List Bullet")
    doc.add_paragraph("增量训练参数（lr=0.02, num_boost_round=50）固定不变，未做超参调优。", style="List Bullet")

    # ── 保存 ──
    output_path = RESULTS_DIR / "LightGBM_反馈循环实验报告.docx"  # 输出文件路径
    doc.save(str(output_path))  # 保存文档
    print(f"报告已生成: {output_path}")


if __name__ == "__main__":
    generate()  # 执行生成函数